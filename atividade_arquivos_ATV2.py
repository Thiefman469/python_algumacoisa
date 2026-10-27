import json
import os
import oracledb
from openpyxl import Workbook, load_workbook
import pandas as pd
os.system("cls")
try:
    conn = oracledb.connect(
        user="rm572510", password="181207", dsn='oracle.fiap.com.br:1521/ORCL')
    # Cursor: onde serão aplicadas as nstruções SQL
    inst_cadastro = conn.cursor()
    inst_consulta = conn.cursor()
    inst_alteracao = conn.cursor()
    inst_exclusao = conn.cursor()
except Exception as e:
    # informa um eventual erro
    print("Erro: ", e)
    conexao = False  # falha
else:
    # Flag para conexao ok
    conexao = True  # ok


# CALCULAR A MÉDIA DAS CPs (desconsiderando a menor nota)

def calcular_media(cp1, cp2, cp3):
    n1 = float(cp1.replace(",", "."))
    n2 = float(cp2.replace(",", "."))
    n3 = float(cp3.replace(",", "."))

    menor = n1

    if n2 < menor:
        menor = n2

    if n3 < menor:
        menor = n3

    media = (n1 + n2 + n3 - menor) / 2

    return round(media, 2)


def gravar_txt(alunos):
    arquivo = open("alunos.txt", "w", encoding="utf-8")

    for aluno in alunos:
        arquivo.write(aluno["rm"] + ";")
        arquivo.write(aluno["nome"] + ";")
        arquivo.write(aluno["cp1"] + ";")
        arquivo.write(aluno["cp2"] + ";")
        arquivo.write(aluno["cp3"] + ";")
        arquivo.write(str(aluno["media_cps"]) + "\n")

    arquivo.close()


# GRAVAR NO ARQUIVO

def gravar_csv(alunos):
    arquivo = open("alunos.csv", "w", encoding="utf-8")

    arquivo.write("RM;Nome;CP1;CP2;CP3;Media CPs\n")

    for aluno in alunos:
        arquivo.write(aluno["rm"] + ";")
        arquivo.write(aluno["nome"] + ";")
        arquivo.write(aluno["cp1"] + ";")
        arquivo.write(aluno["cp2"] + ";")
        arquivo.write(aluno["cp3"] + ";")
        arquivo.write(str(aluno["media_cps"]) + "\n")

    arquivo.close()


# GRAVAR NO ARQUIVO

def gravar_json(alunos):
    arquivo = open("alunos.json", "w", encoding="utf-8")

    json.dump(alunos, arquivo, indent=4, ensure_ascii=False)

    arquivo.close()


# GRAVAR NO ARQUIVO EXCEL

def gravar_excel(alunos):
    # como a lista "alunos" já tem os alunos antigos e os novos,

    arquivo = Workbook()
    planilha = arquivo.active

    planilha["A1"] = "RM"
    planilha["B1"] = "Nome"
    planilha["C1"] = "CP1"
    planilha["D1"] = "CP2"
    planilha["E1"] = "CP3"
    planilha["F1"] = "Media CPs"

    linha = 2

    for aluno in alunos:
        planilha.cell(linha, 1).value = aluno["rm"]
        planilha.cell(linha, 2).value = aluno["nome"]
        planilha.cell(linha, 3).value = aluno["cp1"]
        planilha.cell(linha, 4).value = aluno["cp2"]
        planilha.cell(linha, 5).value = aluno["cp3"]
        planilha.cell(linha, 6).value = aluno["media_cps"]

        linha = linha + 1

    arquivo.save("alunos.xlsx")


# GRAVAR EM TODOS OS ARQUIVOS

def gravar_todos(alunos):
    gravar_txt(alunos)
    gravar_csv(alunos)
    gravar_json(alunos)
    gravar_excel(alunos)

    print()
    print("Cadastro encerrado. Dados gravados nos quatro arquivos.")
    print("------------------------------")


# CARREGAR ALUNOS JÁ CADASTRADOS
# (para não apagar os de execuções anteriores)

def carregar_alunos():
    alunos = []

    if os.path.exists("alunos.txt"):
        arquivo = open("alunos.txt", "r", encoding="utf-8")

        for linha in arquivo:
            linha = linha.strip()

            if linha != "":
                dados = linha.split(";")

                aluno = {
                    "rm": dados[0],
                    "nome": dados[1],
                    "cp1": dados[2],
                    "cp2": dados[3],
                    "cp3": dados[4],
                    "media_cps": float(dados[5])
                }

                alunos.append(aluno)

        arquivo.close()

    return alunos


# MOSTRAR UM ALUNO NA TELA

def mostrar_aluno(rm, nome, cp1, cp2, cp3, media):
    print("RM:", rm, "| Nome:", nome, "| CP1:", cp1,
          "| CP2:", cp2, "| CP3:", cp3, "| Média CPs:", media)


# CADASTRAR ALUNOS

def cadastrar_alunos(alunos):

    print()
    print("------------------------------")
    print("CADASTRO DE ALUNOS")

    rm = input("RM: ")

    while rm != "":

        existe = False

        for aluno in alunos:
            if aluno["rm"] == rm:
                existe = True

        if existe:
            print("RM já cadastrado.")
        else:
            nome = input("Nome: ")
            cp1 = input("CP1: ")
            cp2 = input("CP2: ")
            cp3 = input("CP3: ")

            media_cps = calcular_media(cp1, cp2, cp3)

            aluno = {
                "rm": rm,
                "nome": nome,
                "cp1": cp1,
                "cp2": cp2,
                "cp3": cp3,
                "media_cps": media_cps
            }

            alunos.append(aluno)

            print("Aluno cadastrado.")

    gravar_todos(alunos)


# LER ARQUIVO TXT

def ler_txt():
    print()
    print("------------------------------")

    if os.path.exists("alunos.txt"):
        arquivo = open("alunos.txt", "r", encoding="utf-8")

        print("Conteúdo de alunos.txt")

        for linha in arquivo:
            linha = linha.strip()

            if linha != "":
                dados = linha.split(";")
                mostrar_aluno(dados[0], dados[1], dados[2],
                              dados[3], dados[4], dados[5])

        arquivo.close()
    else:
        print("O arquivo alunos.txt ainda não existe.")
        print("Cadastre pelo menos um aluno antes de tentar lê-lo.")

    print("------------------------------")


# LER ARQUIVO

def ler_csv():
    print()
    print("------------------------------")

    if os.path.exists("alunos.csv"):
        arquivo = open("alunos.csv", "r", encoding="utf-8")

        print("Conteúdo de alunos.csv")

        primeira = True

        for linha in arquivo:
            linha = linha.strip()

            if primeira:
                # a primeira linha é o cabeçalho, então não mostra como aluno
                primeira = False
            elif linha != "":
                dados = linha.split(";")
                mostrar_aluno(dados[0], dados[1], dados[2],
                              dados[3], dados[4], dados[5])

        arquivo.close()
    else:
        print("O arquivo alunos.csv ainda não existe.")
        print("Cadastre pelo menos um aluno antes de tentar lê-lo.")

    print("------------------------------")


# LER ARQUIVO

def ler_json():
    print()
    print("------------------------------")

    if os.path.exists("alunos.json"):
        arquivo = open("alunos.json", "r", encoding="utf-8")

        print("Conteúdo de alunos.json")

        conteudo = json.load(arquivo)

        for aluno in conteudo:
            mostrar_aluno(aluno["rm"], aluno["nome"], aluno["cp1"],
                          aluno["cp2"], aluno["cp3"], aluno["media_cps"])

        arquivo.close()
    else:
        print("O arquivo alunos.json ainda não existe.")
        print("Cadastre pelo menos um aluno antes de tentar lê-lo.")

    print("------------------------------")


# LER ARQUIVO

def ler_excel():
    print()
    print("------------------------------")

    if os.path.exists("alunos.xlsx"):
        arquivo = load_workbook("alunos.xlsx")
        planilha = arquivo.active

        print("Conteúdo de alunos.xlsx")

        # começa na linha 2 porque a linha 1 é o cabeçalho
        linha = 2

        while linha <= planilha.max_row:
            mostrar_aluno(
                planilha.cell(linha, 1).value,
                planilha.cell(linha, 2).value,
                planilha.cell(linha, 3).value,
                planilha.cell(linha, 4).value,
                planilha.cell(linha, 5).value,
                planilha.cell(linha, 6).value
            )

            linha = linha + 1

        arquivo.close()
    else:
        print("O arquivo alunos.xlsx ainda não existe.")
        print("Cadastre pelo menos um aluno antes de tentar lê-lo.")

    print("------------------------------")


# MENU PARA LER ARQUIVO

def ler_arquivo():

    escolha = ""

    while escolha != "0":

        print()
        print("------------------------------")
        print("LER ARQUIVO")
        print("1 - Texto")
        print("2 - CSV")
        print("3 - JSON")
        print("4 - Excel")
        print("0 - Voltar")

        escolha = input("Escolha: ")

        if escolha == "1":
            ler_txt()

        elif escolha == "2":
            ler_csv()

        elif escolha == "3":
            ler_json()

        elif escolha == "4":
            ler_excel()

        elif escolha == "0":
            print("Voltando ao menu principal.")

        else:
            print("Opção inválida. Tente novamente.")


# PESQUISAR ALUNO PELO RM (no arquivo JSON)

def pesquisar_aluno():
    print()
    print("------------------------------")
    print("PESQUISA")

    if os.path.exists("alunos.json"):
        rm = input("Digite o RM do aluno: ")

        arquivo = open("alunos.json", "r", encoding="utf-8")
        conteudo = json.load(arquivo)
        arquivo.close()

        achou = False

        for aluno in conteudo:
            if aluno["rm"] == rm:
                mostrar_aluno(aluno["rm"], aluno["nome"], aluno["cp1"],
                              aluno["cp2"], aluno["cp3"], aluno["media_cps"])
                achou = True

        if not achou:
            print("RM não encontrado.")
    else:
        print("O arquivo alunos.json ainda não existe.")
        print("Cadastre pelo menos um aluno antes de tentar pesquisar.")

    print("------------------------------")

# CRUD SQL


def acesso_sql():
    while conexao:
        os.system("cls")
        print("""
        ------ CRUD - GRAVAÇÃO BD ------
        0 - SAIR
        1 - CADASTRAR ALUNOS
        2 - LISTAR ALUNOS
        3 - EDITAR ALUNOS
        4 - EXCLUIR ALUNOS
        5 - EXCLUIR TODOS OS ALUNOS
        """)
        escolha = int(input("Escolha: "))

        os.system("cls")

        match escolha:
            case 0:
                break
            case 1:  # Cadastrar alunos
                try:
                    print("----- CADASTRAR ALUNOS -----")
                    rm_alunos = input("RM do aluno.....: ")
                    nome = input("Nome.....: ")
                    cp1: int = int(input("CP1.....: "))
                    cp2: int = int(input("CP2.....: "))
                    cp3: int = int(input("CP3.....: "))

                    # montar a instrução SQL de cadastro
                    cadastro = f"""
                    INSERT INTO alunos
                    (rm_aluno, NOME, CP1, CP2, CP3)
                    VALUES ('{rm_alunos}','{nome}', {cp1}, {cp2}, {cp3})
                    """

                    inst_cadastro.execute(cadastro)
                    conn.commit()
                except Exception as e:
                    print("Erro: ", e)
                else:
                    input("\n### DADOS GRAVADOS COM SUCESSO ###")

            case 2:  # listar os alunos
                print("----- LISTAR ALUNOS -----")
                lista_dados = []

                # montar a instrução de consulta
                inst_consulta.execute("SELECT * FROM alunos")
                data = inst_consulta.fetchall()  # joga todos os dados

                if data:  # há dados no objeto data
                    for dt in data:
                        lista_dados.append(dt)
                else:  # Não retornou registros
                    input("Tabela vazia! Aperte ENTER para continuar")

                # ordenar a lista
                lista_dados.sort()

                # Criando um Dataframe (gera uma planilha com os dados)
                dados_df = pd.DataFrame(
                    lista_dados,
                    columns=['rm_alunoss','NOME','CP1','CP2','CP3']
                )

                print(dados_df.to_string(index=False))

                input("\nDados listados! Aperte ENTER para sair")
                
            case 3:  # Editar os alunos
                try:
                    print("----- EDITANDO DADOS DO ALUNO -----")
                    consulta = """
                    SELECT rm_alunos, NOME
                    FROM alunos
                    ORDER BY rm_alunos
                    """

                    inst_consulta.execute(consulta)

                    data = inst_consulta.fetchall()

                    if not data:
                        input("A tabela está vazia! Aperte Enter para continuar")
                        continue

                    dados_df = pd.DataFrame(
                        data,
                        columns=['rm_alunos', 'NOME', 'CP1', 'CP2', 'CP3']
                    )

                    print(dados_df.to_string(index=False))

                    lista_dados = []

                    rm_alunos = int(input("Id: "))

                    consulta = f"""
                    SELECT *
                    FROM alunos
                    WHERE rm_alunos = :rm
                    """

                    inst_consulta.execute(consulta, {"rm": rm_alunos})
                    data = inst_consulta.fetchall()

                    if not data:
                        input("RM inexistente")
                        continue

                    for dt in data:
                        lista_dados.append(dt)

                    novo_rm = input("Novo RM: ")
                    novo_nome = input("Novo Nome: ")
                    nova_nota_CP1 = input("Nova nota CP1: ")
                    nova_nota_CP2 = input("Nova nota CP2: ")
                    nova_nota_CP3 = input("Nova nota CP3: ")

                    # Construção da instrução de edição
                    alteracao = f"""
                    UPDATE alunos
                    SET rm_alunos= :rm, 
                        NOME= :nome,
                        CP1= :cp1,
                        CP2= :cp2, 
                        CP3= :cp3
                    WHERE rm_alunos = :rm
                    """
                    inst_alteracao.execute(alteracao, {
                        "rm": novo_rm,
                        "nome": novo_nome,
                        "cp1": nova_nota_CP1,
                        "cp2": nova_nota_CP2,
                        "cp3": nova_nota_CP3,
                        "rm": rm_alunos
                    })

                    conn.commit()

                except Exception as e:
                    conn.rollback()
                    print("Erro: ", e)
                else:
                    input("\n#### DADOS ATUALIZADOS #####")

            case 4:  # excluir os alunos
                print("----- EXCLUIR ALUNO ----")
                lista_dados = []

                aluno_rm = int(input("RM:"))

                consulta = f"SELECT * FROM alunos WHERE id = {aluno_rm}"
                inst_consulta.execute(consulta)
                data = inst_consulta.fetchall()

                for dt in data:
                    lista_dados.append(dt)

                if len(lista_dados) == 0:
                    print("Não há um aluno cadastrado com este RM")
                else:
                    # Criação de instrução de exclusão
                    exclusao = f"DELETE FROM alunos WHERE id = {aluno_rm}"
                    inst_exclusao.execute(exclusao)
                    conn.commit()
                    input("##### ALUNO EXCLUIDO ")

            case 5:  # excluir todos os alunos
                print("!!!!! EXCLUIR TODOS OS ALUNOS !!!!!")
                confirma = input("Confirma a exclusão? [S]IM OU [NÃO]?")
                if confirma.upper() == "S":
                    try:
                        inst_exclusao.execute("DELETE FROM alunos")
                        conn.commit()
                        print("Todos os alunos foram excluidos")
                    except Exception as e:
                        conn.rollback()
                        print("Erro ao excluir alunos:", e)

                else:
                    print("Exclusão cancelada pelo usuário!")


# PROGRAMA PRINCIPAL
alunos = carregar_alunos()

opcao = ""

while opcao != "0":

    print()
    print("------------------------------")
    print("MENU PRINCIPAL")
    print("0 - SAIR")
    print("1 - Gravar alunos")
    print("2 - Ler arquivo")
    print("3 - Pesquisa")
    print("4 - Gravar dados no BD")

    opcao = input("Escolha: ")

    if opcao == "1":
        cadastrar_alunos(alunos)

    elif opcao == "2":
        ler_arquivo()

    elif opcao == "3":
        pesquisar_aluno()

    elif opcao == "4":
        acesso_sql()

    elif opcao == "0":
        print("Programa encerrado.")

    else:
        print("Opção inválida. Tente novamente.")
